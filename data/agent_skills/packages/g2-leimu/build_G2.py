#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
G2-自己店铺·类目链接明细拆解

用法:
    python build_G2.py --input "生意参谋数据.csv" --output "输出目录" --shop-name "店铺名称"

参数:
    --input      生意参谋商品排行CSV路径（必填）
    --output       输出目录，默认当前工作目录下的 ./outputs
    --shop-name  店铺名称，用于文件名和报告标题
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import argparse
import re
import os

# ==================== 命令行参数 ====================
parser = argparse.ArgumentParser(description="G2-类目链接明细拆解")
parser.add_argument("--input", required=True, help="生意参谋商品排行CSV路径")
parser.add_argument("--output", default=os.path.join(os.getcwd(), "outputs"), help="输出目录，默认当前工作目录下的 ./outputs")
parser.add_argument("--shop-name", default="", help="店铺名称（用于报告标题）")
args = parser.parse_args()

CSV_PATH = args.input
OUTPUT_DIR = args.output
SHOP_NAME = args.shop_name if args.shop_name else os.path.splitext(os.path.basename(CSV_PATH))[0].split("_")[1] if "_" in os.path.basename(CSV_PATH) else "店铺"

base_name = os.path.splitext(os.path.basename(CSV_PATH))[0]
date_match = re.search(r'(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})', base_name)
date_suffix = f"_{date_match.group(1)}至{date_match.group(2)}" if date_match else ""

OUTPUT_PATH = os.path.join(OUTPUT_DIR, f"{SHOP_NAME}_G2_类目链接明细拆解{date_suffix}.xlsx")
os.makedirs(OUTPUT_DIR, exist_ok=True)
print(f"输出路径: {OUTPUT_PATH}")

# ==================== 样式定义 ====================
FILL_TITLE = PatternFill(start_color="2D2006", end_color="2D2006", fill_type="solid")
FILL_HEADER = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
FILL_ALT_ROW = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

FONT_TITLE = Font(name="微软雅黑", size=14, bold=True, color="FFD700")
FONT_HEADER = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="微软雅黑", size=9, color="1A1A1A")
FONT_SUMMARY = Font(name="微软雅黑", size=9, bold=True, color="2D2006")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0')
)

# ==================== 工具函数 ====================
def safe_div(a, b, default=0):
    if b == 0 or pd.isna(b) or pd.isna(a):
        return default
    return a / b

def style_title_row(ws, row, text, cols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = FONT_TITLE
    cell.fill = FILL_TITLE
    cell.alignment = ALIGN_CENTER
    ws.row_dimensions[row].height = 30

def style_summary_row(ws, row, text, cols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = FONT_SUMMARY
    cell.alignment = ALIGN_LEFT
    ws.row_dimensions[row].height = 22

def style_header_row(ws, row, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

def style_data_cell(ws, row, col, value, is_alt=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = FONT_DATA
    if is_alt:
        cell.fill = FILL_ALT_ROW
    else:
        cell.fill = FILL_WHITE
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER

def auto_width(ws, min_width=10, max_width=50):
    for col_idx, col in enumerate(ws.columns, 1):
        max_length = 0
        column = openpyxl.utils.get_column_letter(col_idx)
        for cell in col:
            try:
                if cell.value and not isinstance(cell, openpyxl.cell.cell.MergedCell):
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max(min_width, max_length + 2), max_width)
        ws.column_dimensions[column].width = adjusted_width

def health_status(row):
    """计算健康度评级"""
    refund_rate = row['退款率'] if pd.notna(row['退款率']) else 0
    sales = row['销售额'] if pd.notna(row['销售额']) else 0
    visitors = row['访客数'] if pd.notna(row['访客数']) else 0
    payers = row['支付人数'] if pd.notna(row['支付人数']) else 0
    
    if sales <= 0 and visitors > 0:
        return '⚠️有流量无转化'
    
    s = 100
    if refund_rate > 0.3: s -= 20
    elif refund_rate > 0.2: s -= 10
    if visitors > 0 and payers == 0: s -= 30
    if visitors > 0 and safe_div(sales, visitors, 0) < 1: s -= 15
    
    if s >= 70:
        return '🟢健康'
    elif s >= 50:
        return '🟡中等'
    else:
        return '🟠警告'

def tag_link(row):
    """链接标签：爆款 or 有流量无转化"""
    sales = row['销售额'] if pd.notna(row['销售额']) else 0
    visitors = row['访客数'] if pd.notna(row['访客数']) else 0
    payers = row['支付人数'] if pd.notna(row['支付人数']) else 0
    
    if sales > 50000:
        return '🔥爆款'
    elif visitors > 0 and payers == 0:
        return '⚠️有流量无转化'
    elif sales > 10000:
        return '🔥爆款'
    else:
        return ''

def clean_sheet_name(name):
    """清理Sheet名称中的非法字符"""
    invalid_chars = ['/', '\\', '?', '*', '[', ']', ':']
    for ch in invalid_chars:
        name = name.replace(ch, '-')
    return name[:31]  # Excel Sheet名最大31字符

# ==================== 数据读取与清洗 ====================
print("读取CSV数据...")
df = pd.read_csv(CSV_PATH, encoding='utf-8-sig')

numeric_cols = ['访客数', '销售额', '退款后销售额', 'UV价值', '搜索支付转化率', '支付转化率',
                '平均停留时长(秒)', '收藏率', '加购率', '退款率', '跳失率', '评分',
                '退款金额', '支付人数', '支付件数', '搜索人数', '搜索UV价值', '搜索占比',
                '收藏人数', '加购人数', '加购件数', '客单价', '近7天搜索渠道点击率',
                '总推广花费', '总推广花费占比', '退款后总推广花费占比',
                '关键词推广花费(元)', '关键词推广销售额', '关键词推广投产',
                '万相台花费(元)', '万相台销售额', '万相台投产',
                '精准人群推广花费(元)', '精准人群推广销售额', '精准人群推广投产']

for col in numeric_cols:
    if col in df.columns:
        df[col] = pd.to_numeric(df[col], errors='coerce')

df = df.dropna(subset=['宝贝ID'])

# 提取一级类目
df['一级类目'] = df['商品类目'].str.split('>').str[0]
# 过滤空类目
df = df[df['一级类目'].notna() & (df['一级类目'] != '') & (df['一级类目'] != '-')]
if len(df) == 0:
    raise ValueError("CSV中没有有效的类目数据")

# 修正NaN退款率
df['退款率'] = df.apply(lambda r: safe_div(r['退款金额'], r['销售额'], 0) if pd.isna(r['退款率']) else r['退款率'], axis=1)

# 计算标签
df['标签'] = df.apply(tag_link, axis=1)

# ==================== 创建Excel ====================
print("创建Excel工作簿...")
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ---------- 按类目分组 ----------
cat_groups = df.groupby('一级类目')

# 计算类目索引数据
cat_index_data = []
for cat_name, group in cat_groups:
    links = len(group)
    visitors = group['访客数'].sum()
    sales = group['销售额'].sum()
    refund = group['退款金额'].sum()
    refund_rate = safe_div(refund, sales, 0) * 100
    payers = group['支付人数'].sum()
    search_people = group['搜索人数'].sum()
    promo = group['总推广花费'].sum()
    price = safe_div(sales, payers, 0)
    uv = safe_div(sales, visitors, 0)
    
    # 健康度（用类目平均值判断）
    health = '🟢健康' if refund_rate < 20 and uv > 2 else ('🟡中等' if refund_rate < 40 else '🟠警告')
    
    cat_index_data.append({
        '类目名称': cat_name,
        '链接数': links,
        '访客数': visitors,
        '销售额': sales,
        '退款率': refund_rate,
        '支付人数': payers,
        '搜索人数': search_people,
        '推广花费': promo,
        '客单价': price,
        'UV价值': uv,
        '健康度': health,
    })

cat_df = pd.DataFrame(cat_index_data)
total_sales = cat_df['销售额'].sum()
cat_df['类目占比'] = cat_df['销售额'].apply(lambda x: round(safe_div(x, total_sales, 0) * 100, 1))
cat_df = cat_df.sort_values('销售额', ascending=False)

# ---------- Sheet 1: 类目索引 ----------
ws = wb.create_sheet("1-类目索引")
style_title_row(ws, 1, f"{SHOP_NAME}  类目链接明细索引", 14)

headers = ['序号', '类目名称', '链接数', '访客数', '销售额', '退款率', '支付人数', '搜索人数', '推广花费', '客单价', 'UV价值', '健康度', '类目占比', 'Sheet']
style_header_row(ws, 2, headers)

for i, row in enumerate(cat_df.itertuples(), 3):
    is_alt = (i - 3) % 2 == 1
    sheet_name = clean_sheet_name(row.类目名称)
    vals = [i - 2, row.类目名称, row.链接数, row.访客数, f"¥{row.销售额:,.2f}",
            f"{row.退款率:.1f}%", row.支付人数, row.搜索人数, f"¥{row.推广花费:,.2f}",
            f"{row.客单价:.2f}", f"{row.UV价值:.2f}", row.健康度, f"{row.类目占比}%", sheet_name]
    for col, v in enumerate(vals, 1):
        style_data_cell(ws, i, col, v, is_alt)

auto_width(ws)

# ---------- 各类目明细 Sheet ----------
detail_headers = ['序号', '商品信息', '宝贝ID', '上架时间', '访客数', '销售额', '退款率', '支付人数', '搜索人数', '推广花费', '客单价', 'UV价值', '支付转化率', '收藏率', '加购率', '跳失率', '退款金额', '搜索占比', '直通车花费', '精准人群花费', '标签']

for cat_name, group in cat_groups:
    sheet_name = clean_sheet_name(cat_name)
    
    # 排序：爆款在前，然后按销售额降序
    group = group.copy()
    group['排序权重'] = group['销售额'].apply(lambda x: 1000000 if x > 50000 else (100000 if x > 10000 else x))
    group = group.sort_values('排序权重', ascending=False).drop(columns=['排序权重'])
    
    ws = wb.create_sheet(sheet_name)
    
    # 标题行
    style_title_row(ws, 1, f"{cat_name}  ({len(group)} links)", 21)
    
    # 汇总行
    cat_visitors = group['访客数'].sum()
    cat_sales = group['销售额'].sum()
    cat_refund = group['退款金额'].sum()
    cat_refund_rate = safe_div(cat_refund, cat_sales, 0) * 100
    cat_payers = group['支付人数'].sum()
    cat_promo = group['总推广花费'].sum()
    summary_text = f"Links:{len(group)}  Sales:{cat_sales:.0f}  RefundRate:{cat_refund_rate:.1f}%  Visitors:{cat_visitors}  PayPpl:{cat_payers}  AdCost:{cat_promo:.0f}"
    style_summary_row(ws, 2, summary_text, 21)
    
    # 表头
    style_header_row(ws, 3, detail_headers)
    
    # 数据行
    for i, row in enumerate(group.itertuples(), 4):
        is_alt = (i - 4) % 2 == 1
        
        # 支付转化率格式化
        pay_cr = row.支付转化率 * 100 if pd.notna(row.支付转化率) else 0
        collect_rate = row.收藏率 * 100 if pd.notna(row.收藏率) else 0
        cart_rate = row.加购率 * 100 if pd.notna(row.加购率) else 0
        bounce_rate = row.跳失率 * 100 if pd.notna(row.跳失率) else 0
        search_ratio = row.搜索占比 * 100 if pd.notna(row.搜索占比) else 0
        refund_rate = row.退款率 * 100 if pd.notna(row.退款率) else 0
        
        # 直通车花费 = 关键词推广花费
        ztc_cost = group.loc[row.Index, '关键词推广花费(元)'] if '关键词推广花费(元)' in group.columns and pd.notna(group.loc[row.Index, '关键词推广花费(元)']) else 0
        jz_cost = group.loc[row.Index, '精准人群推广花费(元)'] if '精准人群推广花费(元)' in group.columns and pd.notna(group.loc[row.Index, '精准人群推广花费(元)']) else 0
        
        vals = [
            i - 3,
            row.商品信息[:40] if pd.notna(row.商品信息) else '-',
            str(row.宝贝ID),
            str(row.上架时间) if pd.notna(row.上架时间) else '-',
            row.访客数 if pd.notna(row.访客数) else 0,
            f"¥{row.销售额:,.2f}" if pd.notna(row.销售额) and row.销售额 > 0 else '-',
            f"{refund_rate:.1f}%" if refund_rate > 0 else '-',
            row.支付人数 if pd.notna(row.支付人数) else 0,
            row.搜索人数 if pd.notna(row.搜索人数) else 0,
            f"¥{row.总推广花费:,.2f}" if pd.notna(row.总推广花费) and row.总推广花费 > 0 else 0,
            f"{row.客单价:.2f}" if pd.notna(row.客单价) else 0,
            f"{row.UV价值:.2f}" if pd.notna(row.UV价值) else 0,
            f"{pay_cr:.1f}%" if pay_cr > 0 else 0,
            f"{collect_rate:.1f}%" if collect_rate > 0 else 0,
            f"{cart_rate:.1f}%" if cart_rate > 0 else 0,
            f"{bounce_rate:.1f}%" if bounce_rate > 0 else 0,
            f"¥{row.退款金额:,.2f}" if pd.notna(row.退款金额) and row.退款金额 > 0 else 0,
            f"{search_ratio:.1f}%" if search_ratio > 0 else 0,
            f"¥{ztc_cost:,.2f}" if ztc_cost > 0 else 0,
            f"¥{jz_cost:,.2f}" if jz_cost > 0 else 0,
            row.标签,
        ]
        for col, v in enumerate(vals, 1):
            style_data_cell(ws, i, col, v, is_alt)
    
    auto_width(ws)

# ==================== 保存 ====================
print(f"保存Excel到: {OUTPUT_PATH}")
wb.save(OUTPUT_PATH)
print("完成!")
